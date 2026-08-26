from __future__ import annotations

import html
import io
import re
import zipfile
from copy import deepcopy
from datetime import date
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st


APP_DIR = Path(__file__).parent
TEMPLATE_DIR = APP_DIR / "templates"
SAMPLE_DIR = APP_DIR / "sample_data"
OUTPUT_DIR = APP_DIR / "outputs"
DEFAULT_TEMPLATE_PATH = TEMPLATE_DIR / "template_rps_pste_placeholder.docx"
DEFAULT_TEMPLATE_RELATIVE_PATH = "templates/template_rps_pste_placeholder.docx"
DEFAULT_D3_SAMPLE_PATH = SAMPLE_DIR / "master_rps_d3_pste.xlsx"
DEFAULT_D4_SAMPLE_PATH = SAMPLE_DIR / "master_rps_d4_pste.xlsx"
DEFAULT_SAMPLE_PATH = DEFAULT_D3_SAMPLE_PATH

REQUIRED_SHEETS = [
    "Master_MK",
    "Master_CPL",
    "Master_IK",
    "Master_CPMK",
    "Mapping_MK_CPL",
    "RPS_Pertemuan",
]

OPTIONAL_RPS_SHEETS = ["Short_Silabus", "Referensi", "Evaluasi_RPS", "Asesmen_Mingguan"]
ALLOWED_RPS_SHEETS = set(REQUIRED_SHEETS + OPTIONAL_RPS_SHEETS)


@dataclass(frozen=True)
class SheetSpec:
    name: str
    required_columns: list[str]


SHEET_SPECS = {
    "Master_CPL": SheetSpec("Master_CPL", ["kode_cpl", "deskripsi_cpl"]),
    "Master_IK": SheetSpec("Master_IK", ["kode_ik", "kode_cpl"]),
    "Master_MK": SheetSpec(
        "Master_MK",
        [
            "kode_mk",
            "nama_mk",
            "semester",
        ],
    ),
    "Mapping_MK_CPL": SheetSpec("Mapping_MK_CPL", ["kode_mk", "kode_cpl"]),
    "Master_CPMK": SheetSpec(
        "Master_CPMK",
        ["kode_mk", "kode_cpmk", "deskripsi_cpmk", "kode_ik"],
    ),
    "RPS_Pertemuan": SheetSpec(
        "RPS_Pertemuan",
        [
            "kode_mk",
            "minggu",
            "sub_cpmk",
            "materi",
        ],
    ),
}

RPS_WEEKLY_COLUMNS = [
    "kode_mk",
    "id_penawaran",
    "minggu",
    "sub_cpmk",
    "materi",
    "modalitas",
    "bentuk_pembelajaran",
    "metode",
    "pengalaman_belajar",
    "teknik_asesmen",
    "indikator_penilaian",
    "bobot",
    "referensi",
    "kode_cpmk",
]

RPS_WEEKLY_LABELS = {
    "minggu": "Minggu Ke",
    "sub_cpmk": "Kemampuan akhir yang direncanakan",
    "materi": "Bahan Kajian / Materi Pembelajaran",
    "modalitas": "Modalitas Pembelajaran",
    "bentuk_pembelajaran": "Bentuk Pembelajaran",
    "metode": "Metode Pembelajaran",
    "pengalaman_belajar": "Pengalaman Belajar Mahasiswa",
    "teknik_asesmen": "Teknik Asesmen",
    "indikator_penilaian": "Indikator Penilaian",
    "bobot": "Bobot Penilaian",
    "referensi": "Referensi Mingguan / Sumber Belajar",
}

MODALITAS_OPTIONS = ["Luring", "Daring Sinkron", "Daring Asinkron", "Bauran"]
BENTUK_OPTIONS = [
    "Kuliah",
    "Responsi",
    "Tutorial",
    "Praktikum",
    "Praktik Bengkel",
    "Praktik Lapangan",
    "Proyek",
    "Seminar/Presentasi",
]
METODE_OPTIONS = [
    "Ceramah Interaktif",
    "Diskusi",
    "Case Method",
    "Project Based Learning",
    "Problem Based Learning",
    "Demonstrasi",
    "Praktik Terbimbing",
    "Praktik Mandiri",
]
TEKNIK_ASESMEN_OPTIONS = [
    "Kuis",
    "Tugas",
    "UTS",
    "UAS",
    "Observasi Praktik",
    "Laporan Praktikum",
    "Demonstrasi",
    "Presentasi",
    "Logbook",
    "Portofolio",
    "Rubrik Proyek",
]


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [
        str(col).strip().lower().replace(" ", "_").replace("-", "_") for col in df.columns
    ]
    return df.fillna("")


def apply_sheet_column_aliases(sheet_name: str, df: pd.DataFrame) -> pd.DataFrame:
    aliases = {
        "mata_kuliah": "nama_mk",
        "nama_mata_kuliah": "nama_mk",
        "nama_mk": "nama_mk",
        "rumusan_cpmk": "deskripsi_cpmk",
        "deskripsi_cpmk": "deskripsi_cpmk",
        "rumusan_cpl": "deskripsi_cpl",
        "deskripsi_cpl": "deskripsi_cpl",
        "indikator_kinerja": "deskripsi_ik",
        "deskripsi_ik": "deskripsi_ik",
        "sks": "total_sks",
        "sks_total": "total_sks",
        "total_sks": "total_sks",
        "sks_praktik": "sks_praktek",
        "paket/status": "jenis_mk",
        "jam_teori/minggu": "jam_teori_minggu",
        "jam_praktik/minggu": "jam_praktik_minggu",
        "total_jam/minggu": "total_jam_minggu",
        "bentuk_evaluasi": "bentuk_tes",
        "bentuk_penilaian": "bentuk_tes",
        "jenis_evaluasi": "jenis_tes",
        "jenis_penilaian": "jenis_tes",
        "jenis": "jenis_tes",
        "instrumen": "instrumen_penilaian",
        "rubrik": "rubrik_penilaian",
        "indikator": "indikator_penilaian",
    }
    return df.rename(columns={col: aliases.get(col, col) for col in df.columns}).copy()


def ensure_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    df = df.copy()
    for column in columns:
        if column not in df.columns:
            df[column] = ""
    return df


def normalize_cpl_code(value: Any) -> str:
    raw = str(value).strip().upper()
    if not raw or raw.lower() == "nan":
        return ""
    compact = re.sub(r"[\s_\-]+", "", raw)
    match = re.fullmatch(r"CPL0*(\d+)", compact)
    if match:
        return f"CPL{int(match.group(1))}"
    return compact


def normalize_ik_code(value: Any) -> str:
    raw = str(value).strip().upper()
    if not raw or raw.lower() == "nan":
        return ""
    compact = re.sub(r"\s+", "", raw)
    compact = re.sub(r"^IK[-_]*0*(\d)", r"IK\1", compact)
    return compact


def normalize_kode_mk(value: Any) -> str:
    raw = str(value).strip()
    return "" if raw.lower() == "nan" else raw


def normalize_id_penawaran(value: Any) -> str:
    raw = str(value).strip()
    return "" if raw.lower() == "nan" else raw


def clean_subcpmk_text(text: Any) -> str:
    cleaned = str(text or "").strip()
    if not cleaned or cleaned.lower() == "nan":
        return ""
    cleaned = re.sub(
        r"(?i)\bsub\s*-?\s*cpmk\b\s*(?:\d+(?:[.\-_]\d+)*)?\s*[:;\-.]*\s*",
        "",
        cleaned,
    )
    cleaned = re.sub(r"^\s*[:;\-.]+\s*", "", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def drop_empty_rows(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    cleaned = df.replace(r"^\s*$", pd.NA, regex=True).dropna(how="all")
    return cleaned.fillna("")


def normalize_master_workbook(workbook: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    normalized: dict[str, pd.DataFrame] = {}
    for sheet_name, df in workbook.items():
        df = apply_sheet_column_aliases(sheet_name, drop_empty_rows(df).copy())
        if "kode_mk" in df.columns:
            df["kode_mk"] = df["kode_mk"].map(normalize_kode_mk)
            df = df[df["kode_mk"] != ""].copy()
        if "id_penawaran" in df.columns:
            df["id_penawaran"] = df["id_penawaran"].map(normalize_id_penawaran)
        if sheet_name in ["Master_CPL", "Master_IK", "Mapping_MK_CPL", "Master_CPMK"] and "kode_cpl" in df.columns:
            df["kode_cpl"] = df["kode_cpl"].map(normalize_cpl_code)
            df = df[df["kode_cpl"] != ""].copy()
        if "kode_ik" in df.columns:
            df["kode_ik"] = df["kode_ik"].map(
                lambda value: ", ".join(normalize_ik_code(code) for code in split_codes(value))
            )
        normalized[sheet_name] = df
    if "Short_Silabus" not in normalized:
        normalized["Short_Silabus"] = pd.DataFrame(
            columns=["kode_mk", "id_penawaran", "deskripsi_mk", "bahan_kajian"]
        )
    else:
        normalized["Short_Silabus"] = ensure_columns(
            normalized["Short_Silabus"], ["kode_mk", "id_penawaran", "deskripsi_mk", "bahan_kajian"]
        )
    if "Referensi" not in normalized:
        normalized["Referensi"] = pd.DataFrame(columns=["kode_mk", "id_penawaran", "referensi"])
    else:
        normalized["Referensi"] = ensure_columns(
            normalized["Referensi"], ["kode_mk", "id_penawaran", "referensi"]
        )
    if "Evaluasi_RPS" not in normalized:
        normalized["Evaluasi_RPS"] = pd.DataFrame(
            columns=[
                "kode_mk",
                "id_penawaran",
                "bentuk_tes",
                "jenis_tes",
                "instrumen_penilaian",
                "rubrik_penilaian",
            ]
        )
    else:
        normalized["Evaluasi_RPS"] = ensure_columns(
            normalized["Evaluasi_RPS"],
            [
                "kode_mk",
                "id_penawaran",
                "bentuk_tes",
                "jenis_tes",
                "instrumen_penilaian",
                "rubrik_penilaian",
            ],
        )
    if "Asesmen_Mingguan" not in normalized:
        normalized["Asesmen_Mingguan"] = pd.DataFrame(
            columns=[
                "kode_mk",
                "id_penawaran",
                "minggu",
                "bentuk_tes",
                "jenis_tes",
                "instrumen_penilaian",
                "rubrik_penilaian",
            ]
        )
    else:
        normalized["Asesmen_Mingguan"] = ensure_columns(
            normalized["Asesmen_Mingguan"],
            [
                "kode_mk",
                "id_penawaran",
                "minggu",
                "bentuk_tes",
                "jenis_tes",
                "instrumen_penilaian",
                "rubrik_penilaian",
            ],
        )
    if "Master_MK" in normalized:
        normalized["Master_MK"] = ensure_columns(
            normalized["Master_MK"],
            [
                "kode_mk",
                "id_penawaran",
                "nama_mk",
                "nama_prodi",
                "semester",
                "sks_teori",
                "sks_praktek",
                "jenis_mk",
                "total_sks",
                "jam_teori_minggu",
                "jam_praktik_minggu",
                "total_jam_minggu",
            ],
        )
    if "Master_CPMK" in normalized:
        normalized["Master_CPMK"] = ensure_columns(
            normalized["Master_CPMK"],
            ["kode_mk", "id_penawaran", "kode_cpmk", "deskripsi_cpmk", "kode_ik", "kode_cpl"],
        )
    if "Master_CPL" in normalized:
        normalized["Master_CPL"] = ensure_columns(
            normalized["Master_CPL"], ["kode_cpl", "deskripsi_cpl"]
        )
    if "Master_IK" in normalized:
        normalized["Master_IK"] = ensure_columns(
            normalized["Master_IK"], ["kode_ik", "kode_cpl", "deskripsi_ik"]
        )
    if "RPS_Pertemuan" in normalized:
        normalized["RPS_Pertemuan"] = normalize_rps_pertemuan_columns(
            normalized["RPS_Pertemuan"]
        )
    if "RPS_Pertemuan" in normalized and "Master_CPMK" in normalized:
        cpmk_lookup: dict[tuple[str, str], str] = {}
        for row in normalized["Master_CPMK"].fillna("").to_dict("records"):
            course_key = course_key_from_row(row)
            full_code = str(row.get("kode_cpmk", "")).strip()
            if not full_code:
                continue
            cpmk_lookup[(course_key, full_code)] = full_code
            cpmk_lookup[(course_key, full_code.rsplit("-", 1)[-1])] = full_code

        def canonical_cpmk(row: pd.Series) -> str:
            code = str(row.get("kode_cpmk", "")).strip()
            return cpmk_lookup.get((course_key_from_row(row.to_dict()), code), code)

        normalized["RPS_Pertemuan"]["kode_cpmk"] = normalized["RPS_Pertemuan"].apply(
            canonical_cpmk, axis=1
        )
    if "Master_CPL" in normalized:
        normalized["Master_CPL"] = normalized["Master_CPL"].drop_duplicates(
            subset=["kode_cpl"], keep="first"
        )
    return normalized


def normalize_rps_pertemuan_columns(df: pd.DataFrame) -> pd.DataFrame:
    aliases = {
        "minggu_ke": "minggu",
        "sub_cpmk_/_kemampuan_akhir_yang_direncanakan": "sub_cpmk",
        "kemampuan_akhir_yang_direncanakan": "sub_cpmk",
        "kemampuan_akhir": "sub_cpmk",
        "bahan_kajian_/_materi_pembelajaran": "materi",
        "bahan_kajian": "materi",
        "materi_pembelajaran": "materi",
        "modalitas_pembelajaran": "modalitas",
        "bentuk": "bentuk_pembelajaran",
        "bentuk_pembelajaran": "bentuk_pembelajaran",
        "metode_pembelajaran": "metode",
        "teknik": "teknik_asesmen",
        "indikator": "indikator_penilaian",
        "bobot_penilaian": "bobot",
        "bobot_persen": "bobot",
        "referensi_mingguan_/_sumber_belajar": "referensi",
        "sumber_belajar": "referensi",
    }
    df = df.rename(columns={col: aliases.get(col, col) for col in df.columns}).copy()
    for column in RPS_WEEKLY_COLUMNS:
        if column not in df.columns:
            df[column] = ""
    return df[RPS_WEEKLY_COLUMNS]


def is_practice_course(mk: dict[str, Any]) -> bool:
    sks_praktek = as_int(mk.get("sks_praktek"))
    jenis_mk = str(mk.get("jenis_mk", "")).lower()
    nama_mk = str(mk.get("nama_mk", "")).lower()
    return sks_praktek > 0 or any(
        marker in f"{jenis_mk} {nama_mk}"
        for marker in ["praktik", "praktikum", "bengkel", "proyek", "project"]
    )


def default_bentuk(mk: dict[str, Any], week: int, final_week: int = 17) -> str:
    if week == final_week and is_practice_course(mk):
        return "Proyek"
    if is_practice_course(mk):
        pattern = ["Praktikum", "Praktik Bengkel", "Proyek"]
        return pattern[(week - 1) % len(pattern)]
    pattern = ["Kuliah", "Responsi"]
    return pattern[(week - 1) % len(pattern)]


def default_metode(
    mk: dict[str, Any], week: int, mid_week: int = 9, final_week: int = 17
) -> str:
    if week == mid_week:
        return "Problem Based Learning"
    if week == final_week and is_practice_course(mk):
        return "Project Based Learning"
    if is_practice_course(mk):
        pattern = [
            "Demonstrasi",
            "Praktik Terbimbing",
            "Praktik Mandiri",
            "Project Based Learning",
        ]
        return pattern[(week - 1) % len(pattern)]
    pattern = ["Ceramah Interaktif", "Diskusi", "Case Method"]
    return pattern[(week - 1) % len(pattern)]


def default_weekly_assessment(
    mk: dict[str, Any], week: int, mid_week: int = 9, final_week: int = 17
) -> tuple[str, float]:
    if week == mid_week:
        return "UTS", 20.0
    if week == final_week:
        return ("Rubrik Proyek", 40.0) if is_practice_course(mk) else ("UAS", 40.0)
    if week in [4, 12]:
        return "Tugas", 10.0
    if week in [2, 6, 10, 14]:
        return "Kuis", 5.0
    return "", 0.0


def normalize_weekly_df(
    weekly_records: list[dict[str, Any]], mk: dict[str, Any], cpmk_records: list[dict[str, Any]]
) -> pd.DataFrame:
    existing_df = normalize_rps_pertemuan_columns(pd.DataFrame(weekly_records))
    existing_by_week = {
        as_int(row.get("minggu")): row for row in existing_df.fillna("").to_dict("records")
        if as_int(row.get("minggu")) > 0
    }
    cpmk_codes = [
        str(row.get("kode_cpmk", "")).strip()
        for row in cpmk_records
        if str(row.get("kode_cpmk", "")).strip()
    ]
    rows: list[dict[str, Any]] = []
    final_week = max(existing_by_week) if existing_by_week else 17
    mid_week = 9
    for week in range(1, final_week + 1):
        row = {column: "" for column in RPS_WEEKLY_COLUMNS}
        row["kode_mk"] = mk.get("kode_mk", "")
        row["id_penawaran"] = mk.get("id_penawaran", "")
        row["minggu"] = week
        row["modalitas"] = "Luring"
        row["bentuk_pembelajaran"] = default_bentuk(mk, week, final_week)
        row["metode"] = default_metode(mk, week, mid_week, final_week)
        row["kode_cpmk"] = cpmk_codes[(week - 1) % len(cpmk_codes)] if cpmk_codes else ""
        row["pengalaman_belajar"] = "Mahasiswa mengikuti aktivitas pembelajaran dan menyelesaikan tugas terarah."
        assessment, weight = default_weekly_assessment(mk, week, mid_week, final_week)
        row["teknik_asesmen"] = assessment
        row["bobot"] = weight
        if week == mid_week:
            row["sub_cpmk"] = "Evaluasi tengah semester"
            row["materi"] = "UTS"
            row["indikator_penilaian"] = "Ketepatan penyelesaian soal UTS"
        elif week == final_week:
            row["sub_cpmk"] = "Evaluasi akhir atau proyek akhir semester"
            row["materi"] = "UAS / proyek akhir semester"
            row["indikator_penilaian"] = "Ketercapaian capaian pembelajaran akhir"
        if week in existing_by_week:
            for key, value in existing_by_week[week].items():
                if str(value).strip() != "":
                    row[key] = value
        if row["modalitas"] not in MODALITAS_OPTIONS:
            row["modalitas"] = "Luring"
        if row["bentuk_pembelajaran"] not in BENTUK_OPTIONS:
            row["bentuk_pembelajaran"] = default_bentuk(mk, week, final_week)
        if row["metode"] not in METODE_OPTIONS:
            row["metode"] = default_metode(mk, week, mid_week, final_week)
        rows.append(row)
    return pd.DataFrame(rows, columns=RPS_WEEKLY_COLUMNS)


@st.cache_data(show_spinner=False)
def load_master_excel(file_bytes: bytes) -> dict[str, pd.DataFrame]:
    excel = pd.ExcelFile(io.BytesIO(file_bytes), engine="openpyxl")
    selected_sheets = [name for name in excel.sheet_names if name in ALLOWED_RPS_SHEETS]
    sheets = {
        name: pd.read_excel(excel, sheet_name=name)
        for name in selected_sheets
    }
    workbook = {
        name: normalize_columns(df)
        for name, df in sheets.items()
        if name in ALLOWED_RPS_SHEETS
    }
    return normalize_master_workbook(workbook)


def validate_workbook_schema(workbook: dict[str, pd.DataFrame]) -> list[str]:
    errors: list[str] = []
    for sheet_name, spec in SHEET_SPECS.items():
        if sheet_name not in workbook:
            errors.append(f"Sheet `{sheet_name}` belum tersedia.")
            continue
        missing = [col for col in spec.required_columns if col not in workbook[sheet_name].columns]
        if missing:
            errors.append(
                f"Sheet `{sheet_name}` kurang kolom: {', '.join(f'`{col}`' for col in missing)}."
            )
    return errors


def split_codes(value: Any) -> list[str]:
    if value is None or str(value).strip() == "":
        return []
    return [item.strip() for item in re.split(r"[,;/\n]+", str(value)) if item.strip()]


def as_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def as_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(str(value).replace("%", "").strip())
    except (TypeError, ValueError):
        return default


def records_to_editor(records: list[dict[str, Any]], columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(records, columns=columns).fillna("")


def course_key_from_row(row: dict[str, Any]) -> str:
    return normalize_id_penawaran(row.get("id_penawaran", "")) or normalize_kode_mk(
        row.get("kode_mk", "")
    )


def dataframe_course_keys(df: pd.DataFrame) -> pd.Series:
    codes = df.get("kode_mk", pd.Series("", index=df.index)).map(normalize_kode_mk)
    offers = df.get("id_penawaran", pd.Series("", index=df.index)).map(
        normalize_id_penawaran
    )
    return offers.where(offers != "", codes)


def filter_by_course(
    df: pd.DataFrame, kode_mk: str, id_penawaran: str = ""
) -> pd.DataFrame:
    if df.empty or "kode_mk" not in df.columns:
        return df.iloc[0:0].copy()
    selected_offer = normalize_id_penawaran(id_penawaran)
    if selected_offer and "id_penawaran" in df.columns:
        normalized_offers = df["id_penawaran"].map(normalize_id_penawaran)
        if (normalized_offers == selected_offer).any():
            return df[normalized_offers == selected_offer].copy()
        if (normalized_offers != "").any():
            return df.iloc[0:0].copy()
    selected_code = normalize_kode_mk(kode_mk)
    normalized_codes = df["kode_mk"].map(normalize_kode_mk)
    return df[normalized_codes == selected_code].copy()


def filter_by_code(df: pd.DataFrame, kode_mk: str) -> pd.DataFrame:
    """Backward-compatible filter for workbooks without id_penawaran."""
    return filter_by_course(df, kode_mk)


def unique_values(values: list[Any]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value).strip()
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def explode_cpmk_ik(cpmk_df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for row in cpmk_df.fillna("").to_dict("records"):
        ik_codes = [normalize_ik_code(code) for code in split_codes(row.get("kode_ik", ""))]
        for kode_ik in ik_codes:
            if kode_ik:
                exploded = row.copy()
                exploded["kode_ik"] = kode_ik
                rows.append(exploded)
    return pd.DataFrame(rows)


def build_course_payload(workbook: dict[str, pd.DataFrame], course_key: str) -> dict[str, Any]:
    mk_df = workbook["Master_MK"]
    selected_key = str(course_key).strip()
    offer_matches = mk_df["id_penawaran"].map(normalize_id_penawaran) == selected_key
    if offer_matches.any():
        mk_row = mk_df[offer_matches].iloc[0].to_dict()
    else:
        code_matches = mk_df["kode_mk"].map(normalize_kode_mk) == normalize_kode_mk(selected_key)
        if not code_matches.any():
            raise ValueError(f"Mata kuliah `{selected_key}` tidak ditemukan di Master_MK.")
        mk_row = mk_df[code_matches].iloc[0].to_dict()
    selected_code = normalize_kode_mk(mk_row.get("kode_mk", ""))
    selected_offer = normalize_id_penawaran(mk_row.get("id_penawaran", ""))
    warnings: list[str] = []

    cpmk_df = filter_by_course(workbook["Master_CPMK"], selected_code, selected_offer)
    cpmk_records = cpmk_df.to_dict("records")
    cpl_codes = unique_values(cpmk_df["kode_cpl"].map(normalize_cpl_code).tolist())
    cpl_df = workbook["Master_CPL"]
    cpl_records = cpl_df[cpl_df["kode_cpl"].astype(str).isin(cpl_codes)].drop_duplicates(
        subset=["kode_cpl"], keep="first"
    ).to_dict("records")
    missing_cpl_from_cpmk = sorted(set(cpl_codes) - set(cpl_df["kode_cpl"].astype(str)))
    for kode_cpl in missing_cpl_from_cpmk:
        warnings.append(f"Kode CPL `{kode_cpl}` dari Master_CPMK tidak ditemukan di Master_CPL.")
    cpmk_ik_df = explode_cpmk_ik(cpmk_df)
    cpmk_ik_codes = unique_values(
        cpmk_ik_df["kode_ik"].astype(str).tolist() if not cpmk_ik_df.empty else []
    )
    ik_df = workbook["Master_IK"]
    ik_records = ik_df[ik_df["kode_ik"].astype(str).isin(cpmk_ik_codes)].drop_duplicates(
        subset=["kode_ik"], keep="first"
    ).to_dict("records")
    found_ik = {str(row.get("kode_ik", "")) for row in ik_records}
    if not cpmk_ik_df.empty:
        fallback_ik_records = (
            cpmk_ik_df[~cpmk_ik_df["kode_ik"].astype(str).isin(found_ik)]
            .drop_duplicates(subset=["kode_ik"], keep="first")
            .to_dict("records")
        )
        for row in fallback_ik_records:
            ik_records.append(
                {
                    "kode_ik": row.get("kode_ik", ""),
                    "kode_cpl": normalize_cpl_code(row.get("kode_cpl", "")),
                    "deskripsi_ik": row.get("deskripsi_ik", ""),
                }
            )
    missing_ik = sorted(set(cpmk_ik_codes) - set(ik_df["kode_ik"].astype(str)))
    for kode_ik in missing_ik:
        warnings.append(f"Kode IK `{kode_ik}` dari Master_CPMK tidak ditemukan di Master_IK.")
    cpl_codes = unique_values(
        cpl_codes + [normalize_cpl_code(row.get("kode_cpl", "")) for row in ik_records]
    )
    cpl_records = cpl_df[cpl_df["kode_cpl"].astype(str).isin(cpl_codes)].drop_duplicates(
        subset=["kode_cpl"], keep="first"
    ).to_dict("records")
    ik_cpl_codes = unique_values([str(row.get("kode_cpl", "")) for row in ik_records])
    missing_cpl_from_ik = sorted(set(ik_cpl_codes) - set(cpl_df["kode_cpl"].astype(str)))
    for kode_cpl in missing_cpl_from_ik:
        warnings.append(f"Kode CPL `{kode_cpl}` dari Master_IK tidak ditemukan di Master_CPL.")

    silabus_df = filter_by_course(workbook["Short_Silabus"], selected_code, selected_offer)
    silabus = silabus_df.iloc[0].to_dict() if not silabus_df.empty else {}

    weekly_records = filter_by_course(workbook["RPS_Pertemuan"], selected_code, selected_offer).sort_values(
        by="minggu", key=lambda s: pd.to_numeric(s, errors="coerce")
    ).to_dict("records")
    if not weekly_records:
        warnings.append("Data RPS pertemuan untuk mata kuliah ini belum tersedia.")
    reference_records = filter_by_course(workbook["Referensi"], selected_code, selected_offer).to_dict("records")
    evaluation_records = filter_by_course(workbook["Evaluasi_RPS"], selected_code, selected_offer).to_dict("records")
    weekly_assessment_records = filter_by_course(workbook["Asesmen_Mingguan"], selected_code, selected_offer).to_dict("records")

    sks_teori = as_int(mk_row.get("sks_teori"))
    sks_praktek = as_int(mk_row.get("sks_praktek"))
    total_sks = as_int(mk_row.get("total_sks"), sks_teori + sks_praktek)

    return {
        "mk": mk_row,
        "course_key": course_key_from_row(mk_row),
        "cpl": cpl_records,
        "ik": ik_records,
        "silabus": silabus,
        "cpmk": cpmk_records,
        "weekly": normalize_weekly_df(weekly_records, mk_row, cpmk_records).to_dict("records"),
        "references": reference_records,
        "evaluations": evaluation_records,
        "weekly_assessments": weekly_assessment_records,
        "total_sks": total_sks,
        "warnings": warnings,
    }


def bullet_text(records: list[dict[str, Any]], code_key: str, text_key: str) -> str:
    if not records:
        return "-"
    lines = []
    for row in records:
        code = str(row.get(code_key, "")).strip()
        text = str(row.get(text_key, "")).strip()
        lines.append(f"{code} - {text}" if code else text)
    return "\n".join(lines)


def reference_text(reference_df: pd.DataFrame) -> str:
    if reference_df.empty:
        return "-"
    return "\n".join(
        f"{idx}. {row.get('referensi', '')}"
        for idx, row in enumerate(reference_df.to_dict("records"), start=1)
    )


def cpmk_text(cpmk_df: pd.DataFrame) -> str:
    records = cpmk_df.fillna("").to_dict("records")
    if not records:
        return "-"
    lines = []
    for row in records:
        lines.append(
            " | ".join(
                [
                    str(row.get("kode_cpmk", "")).strip(),
                    str(row.get("deskripsi_cpmk", "")).strip(),
                    f"IK: {str(row.get('kode_ik', '')).strip()}",
                ]
            )
        )
    return "\n".join(lines)


def weekly_plan_text(weekly_df: pd.DataFrame) -> str:
    records = weekly_df.fillna("").to_dict("records")
    if not records:
        return "-"
    lines = []
    for row in records:
        lines.append(
            "\n".join(
                [
                    f"Minggu {row.get('minggu', '')}",
                    f"Kemampuan akhir yang direncanakan: {clean_subcpmk_text(row.get('sub_cpmk', ''))}",
                    f"Bahan Kajian / Materi: {row.get('materi', '')}",
                    f"Modalitas: {row.get('modalitas', '')}",
                    f"Bentuk Pembelajaran: {row.get('bentuk_pembelajaran', '')}",
                    f"Metode Pembelajaran: {row.get('metode', '')}",
                    f"Pengalaman Belajar Mahasiswa: {row.get('pengalaman_belajar', '')}",
                    f"Teknik Asesmen: {row.get('teknik_asesmen', '')}",
                    f"Indikator: {row.get('indikator_penilaian', '')}",
                    f"Bobot: {row.get('bobot', '')}%",
                    f"Referensi Mingguan / Sumber Belajar: {row.get('referensi', '')}",
                ]
            )
        )
    return "\n\n".join(lines)


def assessment_text(weekly_df: pd.DataFrame) -> str:
    records = [
        row
        for row in weekly_df.fillna("").to_dict("records")
        if str(row.get("teknik_asesmen", "")).strip()
        or str(row.get("indikator_penilaian", "")).strip()
        or as_float(row.get("bobot")) > 0
    ]
    if not records:
        return "-"
    return "\n".join(
        " | ".join(
            [
                f"Minggu {row.get('minggu', '')}",
                clean_subcpmk_text(row.get("sub_cpmk", "")),
                str(row.get("teknik_asesmen", "")).strip(),
                str(row.get("indikator_penilaian", "")).strip(),
                f"{row.get('bobot', '')}%",
            ]
        )
        for row in records
    )


def make_context(
    payload: dict[str, Any],
    lecturer_name: str,
    description: str,
    cpmk_df: pd.DataFrame,
    weekly_df: pd.DataFrame,
    reference_df: pd.DataFrame,
) -> dict[str, Any]:
    mk = payload["mk"]
    references = reference_df.fillna("").to_dict("records")
    pustaka_utama = references[0].get("referensi", "") if references else ""
    pustaka_pendukung = "\n".join(
        str(row.get("referensi", "")).strip()
        for row in references[1:]
        if str(row.get("referensi", "")).strip()
    )
    ka_prodi = (
        mk.get("nama_ka_prodi")
        or mk.get("ka_prodi")
        or mk.get("kaprodi")
        or mk.get("nama_kaprodi")
        or mk.get("nama_kakel_bidang_keahlian")
        or ""
    )
    evaluation_records = payload.get("evaluations", [])
    first_evaluation = evaluation_records[0] if evaluation_records else {}
    context = {
        "NAMA_PRODI": mk.get("nama_prodi", ""),
        "KODE_MK": mk.get("kode_mk", ""),
        "NAMA_MK": mk.get("nama_mk", ""),
        "SEMESTER": mk.get("semester", ""),
        "SKS_TEORI": mk.get("sks_teori", ""),
        "SKS_PRAKTEK": mk.get("sks_praktek", ""),
        "TOTAL_SKS": payload["total_sks"],
        "DOSEN_PENGAMPU": lecturer_name,
        "DESKRIPSI_MK": description,
        "CPL_DIBEBANKAN": bullet_text(payload["cpl"], "kode_cpl", "deskripsi_cpl"),
        "IK_TERKAIT": bullet_text(payload["ik"], "kode_ik", "deskripsi_ik"),
        "BAHAN_KAJIAN": payload["silabus"].get("bahan_kajian", ""),
        "CPMK_TEXT": cpmk_text(cpmk_df),
        "RENCANA_MINGGUAN_TEXT": weekly_plan_text(weekly_df),
        "ASESMEN_TEXT": assessment_text(weekly_df),
        "REFERENSI": reference_text(reference_df),
    }
    context.update(
        {
            "nama_prodi": context["NAMA_PRODI"],
            "kode_mk": context["KODE_MK"],
            "nama_mk": context["NAMA_MK"],
            "semester": context["SEMESTER"],
            "sks_teori": context["SKS_TEORI"],
            "sks_praktek": context["SKS_PRAKTEK"],
            "total_sks": context["TOTAL_SKS"],
            "dosen_pengampu": context["DOSEN_PENGAMPU"],
            "deskripsi_mk": context["DESKRIPSI_MK"],
            "cpl_text": context["CPL_DIBEBANKAN"],
            "ik_text": context["IK_TERKAIT"],
            "bahan_kajian": context["BAHAN_KAJIAN"],
            "cpmk_text": context["CPMK_TEXT"],
            "rencana_mingguan_text": context["RENCANA_MINGGUAN_TEXT"],
            "asesmen_text": context["ASESMEN_TEXT"],
            "referensi_text": context["REFERENSI"],
            "sks_total": context["TOTAL_SKS"],
            "tgl_penyusunan": date.today().strftime("%d-%m-%Y"),
            "nama_kakel_bidang_keahlian": ka_prodi,
            "nama_ka_prodi": ka_prodi,
            "mata_kuliah_syarat": mk.get("mata_kuliah_syarat", ""),
            "pustaka_utama": pustaka_utama,
            "pustaka_pendukung": pustaka_pendukung,
            "rumus_nilai_akhir": "Nilai Akhir = jumlah seluruh nilai komponen penilaian x bobot masing-masing.",
            "bentuk_tes": first_evaluation.get("bentuk_tes", ""),
            "jenis_tes": first_evaluation.get("jenis_tes", ""),
            "instrumen_penilaian": first_evaluation.get("instrumen_penilaian", ""),
            "rubrik_penilaian": first_evaluation.get("rubrik_penilaian", ""),
            "kisi_kisi_instrumen": first_evaluation.get("kisi_kisi_instrumen", ""),
        }
    )
    return context


def validate_rps(
    payload: dict[str, Any],
    cpmk_df: pd.DataFrame,
    weekly_df: pd.DataFrame,
) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    ik_df = pd.DataFrame(payload["ik"])
    valid_ik = set(ik_df["kode_ik"].map(normalize_ik_code)) if not ik_df.empty else set()
    ik_to_cpl = {
        normalize_ik_code(row["kode_ik"]): normalize_cpl_code(row["kode_cpl"])
        for row in ik_df.to_dict("records")
        if str(row.get("kode_ik", "")).strip()
    }
    cpl_codes = {normalize_cpl_code(row.get("kode_cpl", "")) for row in payload["cpl"]}

    for warning in payload.get("warnings", []):
        rows.append(
            {
                "status": "Warning",
                "aturan": "Kode master tidak ditemukan",
                "detail": warning,
            }
        )

    used_ik: set[str] = set()
    used_cpl: set[str] = set()
    for row in cpmk_df.fillna("").to_dict("records"):
        kode_cpmk = str(row.get("kode_cpmk", "")).strip() or "(CPMK tanpa kode)"
        ik_codes = [normalize_ik_code(code) for code in split_codes(row.get("kode_ik", ""))]
        if not ik_codes:
            rows.append(
                {
                    "status": "Error",
                    "aturan": "Setiap CPMK harus punya IK",
                    "detail": f"{kode_cpmk} belum memiliki kode IK.",
                }
            )
        for kode_ik in ik_codes:
            used_ik.add(kode_ik)
            if kode_ik not in valid_ik:
                rows.append(
                    {
                        "status": "Warning",
                        "aturan": "Kode IK tidak ditemukan di master",
                        "detail": f"{kode_cpmk} memakai IK `{kode_ik}` yang tidak ditemukan di Master_IK untuk mata kuliah ini.",
                    }
                )
                continue
            kode_cpl = ik_to_cpl.get(kode_ik, "")
            if not kode_cpl:
                rows.append(
                    {
                        "status": "Error",
                        "aturan": "Setiap IK harus punya CPL",
                        "detail": f"IK `{kode_ik}` belum dipetakan ke CPL.",
                    }
                )
            elif kode_cpl not in cpl_codes:
                rows.append(
                    {
                        "status": "Warning",
                        "aturan": "Kode CPL IK tidak termasuk CPL dibebankan",
                        "detail": f"IK `{kode_ik}` terhubung ke `{kode_cpl}`, tetapi CPL tersebut tidak muncul pada Master_CPMK untuk mata kuliah ini.",
                    }
                )
                used_cpl.add(kode_cpl)
            else:
                used_cpl.add(kode_cpl)

    missing_cpl = sorted(cpl_codes - used_cpl)
    for kode_cpl in missing_cpl:
        rows.append(
            {
                "status": "Warning",
                "aturan": "Setiap CPL dibebankan muncul minimal pada satu CPMK",
                "detail": f"CPL `{kode_cpl}` belum muncul pada CPMK.",
            }
        )

    total_weight = weekly_df.get("bobot", pd.Series(dtype=float)).map(as_float).sum()
    if round(total_weight, 2) != 100:
        rows.append(
            {
                "status": "Warning",
                "aturan": "Total bobot penilaian harus 100%",
                "detail": f"Total bobot saat ini {total_weight:g}%.",
            }
        )

    final_week = max(weekly_df.get("minggu", pd.Series([17])).map(as_int), default=17)
    mid_week = 9
    for row in weekly_df.fillna("").to_dict("records"):
        week = row.get("minggu", "")
        sub_cpmk = str(row.get("sub_cpmk", "")).strip()
        assessment = str(row.get("teknik_asesmen", "")).strip()
        if not sub_cpmk:
            rows.append(
                {
                    "status": "Warning",
                    "aturan": "Setiap pertemuan sebaiknya memiliki kemampuan akhir",
                    "detail": f"Pertemuan minggu {week} belum memiliki kemampuan akhir yang direncanakan.",
                }
            )
        if sub_cpmk and not assessment:
            rows.append(
                {
                    "status": "Warning",
                    "aturan": "Kemampuan akhir perlu asesmen",
                    "detail": f"Pertemuan minggu {week} memiliki kemampuan akhir tetapi belum ada teknik asesmen.",
                }
            )
        week_num = as_int(week)
        materi = str(row.get("materi", "")).upper()
        teknik = str(row.get("teknik_asesmen", "")).upper()
        if week_num == mid_week and "UTS" not in f"{materi} {teknik}":
            rows.append(
                {
                    "status": "Warning",
                    "aturan": f"Minggu {mid_week} wajib UTS",
                    "detail": f"Pertemuan minggu {mid_week} belum ditandai sebagai UTS.",
                }
            )
        if week_num == final_week and "UAS" not in f"{materi} {teknik}":
            rows.append(
                {
                    "status": "Warning",
                    "aturan": f"Minggu {final_week} wajib UAS",
                    "detail": f"Pertemuan minggu {final_week} belum ditandai sebagai UAS.",
                }
            )

    if not rows:
        rows.append(
            {
                "status": "OK",
                "aturan": "Semua validasi terpenuhi",
                "detail": "RPS siap diekspor.",
            }
        )
    return pd.DataFrame(rows)


def validate_master_data(workbook: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    master_mk = workbook.get("Master_MK", pd.DataFrame())
    rps_df = workbook.get("RPS_Pertemuan", pd.DataFrame())
    cpmk_df = workbook.get("Master_CPMK", pd.DataFrame())
    ik_df = workbook.get("Master_IK", pd.DataFrame())

    mk_keys = set(dataframe_course_keys(master_mk))

    if not rps_df.empty:
        rps_keys = set(dataframe_course_keys(rps_df))
        for course_key in sorted(rps_keys - mk_keys):
            rows.append(
                {
                    "status": "Warning",
                    "aturan": "Penawaran RPS harus ada di Master_MK",
                    "detail": f"Penawaran `{course_key}` ada di RPS_Pertemuan tetapi tidak ada di Master_MK.",
                }
            )

        valid_cpmk_pairs = {
            (course_key_from_row(row), str(row.get("kode_cpmk", "")).strip())
            for row in cpmk_df.fillna("").to_dict("records")
        }
        for row in rps_df.fillna("").to_dict("records"):
            pair = (course_key_from_row(row), str(row.get("kode_cpmk", "")).strip())
            if pair[1] and pair not in valid_cpmk_pairs:
                rows.append(
                    {
                        "status": "Warning",
                        "aturan": "Kode CPMK RPS harus ada di Master_CPMK",
                        "detail": f"`{pair[1]}` pada `{pair[0]}` tidak ditemukan di Master_CPMK.",
                    }
                )

        grouped_rps = rps_df.assign(_course_key=dataframe_course_keys(rps_df))
        for course_key, group in grouped_rps.groupby("_course_key", dropna=False):
            total_weight = group.get("bobot", pd.Series(dtype=float)).map(as_float).sum()
            if round(total_weight, 2) != 100:
                rows.append(
                    {
                        "status": "Warning",
                        "aturan": "Total bobot RPS per MK harus 100%",
                        "detail": f"`{course_key}` memiliki total bobot {total_weight:g}%.",
                    }
                )
            week_map = {
                as_int(row.get("minggu")): row for row in group.fillna("").to_dict("records")
            }
            final_week = max(week_map, default=17)
            for week, label in [(9, "UTS"), (final_week, "UAS")]:
                row = week_map.get(week, {})
                text = f"{row.get('materi', '')} {row.get('teknik_asesmen', '')}".upper()
                if label not in text:
                    rows.append(
                        {
                            "status": "Warning",
                            "aturan": f"Minggu {week} wajib {label}",
                            "detail": f"`{course_key}` belum menandai minggu {week} sebagai {label}.",
                        }
                    )
            for row in group.fillna("").to_dict("records"):
                week = as_int(row.get("minggu"))
                if week in [9, final_week]:
                    continue
                if not str(row.get("materi", "")).strip() or not str(row.get("sub_cpmk", "")).strip():
                    rows.append(
                        {
                            "status": "Warning",
                            "aturan": "Minggu reguler harus punya materi dan kemampuan akhir",
                            "detail": f"`{course_key}` minggu {week} belum lengkap.",
                        }
                    )

    for row in cpmk_df.fillna("").to_dict("records"):
        kode_cpmk = str(row.get("kode_cpmk", "")).strip()
        if not split_codes(row.get("kode_ik", "")):
            rows.append(
                {
                    "status": "Warning",
                    "aturan": "Setiap CPMK punya kode IK",
                    "detail": f"`{kode_cpmk}` belum memiliki kode IK.",
                }
            )

    for row in ik_df.fillna("").to_dict("records"):
        kode_ik = str(row.get("kode_ik", "")).strip()
        if kode_ik and not normalize_cpl_code(row.get("kode_cpl", "")):
            rows.append(
                {
                    "status": "Warning",
                    "aturan": "Setiap IK punya kode CPL",
                    "detail": f"`{kode_ik}` belum memiliki kode CPL.",
                }
            )

    if not rows:
        rows.append(
            {
                "status": "OK",
                "aturan": "Validasi data master",
                "detail": "Tidak ada mismatch utama pada workbook.",
            }
        )
    return pd.DataFrame(rows)


def read_docx_document_xml(template_bytes: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(template_bytes)) as docx:
            return docx.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception as exc:
        raise ValueError(f"Template Word tidak dapat dibaca sebagai file .docx: {exc}") from exc


def validate_template_loops(template_bytes: bytes) -> None:
    xml = read_docx_document_xml(template_bytes)
    visible_text = html.unescape(re.sub(r"<[^>]+>", "", xml))
    for_count = max(
        len(re.findall(r"{%-?\s*(?:tr\s+)?for\b", xml)),
        len(re.findall(r"{%-?\s*(?:tr\s+)?for\b", visible_text)),
    )
    endfor_count = max(
        len(re.findall(r"{%-?\s*(?:tr\s+)?endfor\s*-?%}", xml)),
        len(re.findall(r"{%-?\s*(?:tr\s+)?endfor\s*-?%}", visible_text)),
    )
    if for_count != endfor_count:
        raise ValueError(
            "Template Word memiliki tag loop Jinja yang tidak seimbang: "
            f"ditemukan {for_count} tag for dan {endfor_count} tag endfor. "
            "Untuk versi stabil ini, gunakan placeholder sederhana seperti "
            "{{ nama_mk }}, {{ cpl_text }}, {{ cpmk_text }}, dan "
            "{{ rencana_mingguan_text }} tanpa {% for %} / {% endfor %}."
        )
    if for_count or endfor_count:
        raise ValueError(
            "Template Word masih memakai tag loop Jinja. Untuk versi stabil ini, "
            "hapus {% for %} / {% endfor %} dan gunakan placeholder teks multiline "
            "seperti {{ cpl_text }}, {{ ik_text }}, {{ cpmk_text }}, "
            "{{ rencana_mingguan_text }}, {{ asesmen_text }}, dan {{ referensi_text }}."
        )


def apply_table_style_safe(doc, table, preferred_styles=None):
    if preferred_styles is None:
        preferred_styles = [
            "Table Grid",
            "TableGrid",
            "Light Grid",
            "Light List",
            "Grid Table 1 Light",
        ]
    try:
        from docx.enum.style import WD_STYLE_TYPE

        table_styles = [
            style
            for style in doc.styles
            if getattr(style, "type", None) == WD_STYLE_TYPE.TABLE
        ]
        styles_by_name = {style.name: style for style in table_styles}
        styles_by_id = {getattr(style, "style_id", ""): style for style in table_styles}
        for preferred_style in preferred_styles:
            style = styles_by_name.get(preferred_style) or styles_by_id.get(preferred_style)
            if style is None:
                continue
            try:
                table.style = style
                return style.name
            except Exception:
                continue
    except Exception:
        return None
    return None


def replace_text_placeholders(text: str, context: dict[str, Any]) -> str:
    rendered = text
    for key, value in context.items():
        value_text = "" if value is None else str(value)
        placeholders = [
            "{{" + key + "}}",
            "{{ " + key + " }}",
        ]
        for placeholder in placeholders:
            rendered = rendered.replace(placeholder, value_text)
    return rendered


def set_cell_text(cell, text: Any) -> None:
    value = "" if text is None else str(text)
    if cell.paragraphs:
        paragraph = cell.paragraphs[0]
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = value
        else:
            paragraph.add_run(value)
        for paragraph in cell.paragraphs[1:]:
            for run in paragraph.runs:
                run.text = ""
    else:
        cell.text = value


def row_text(row) -> str:
    return " ".join(cell.text for cell in row.cells)


def replace_placeholders_in_row(row, context: dict[str, Any]) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            replace_placeholders_in_paragraph(paragraph, context)


def remove_table_row(row) -> None:
    tr = row._tr
    tr.getparent().remove(tr)


def fill_repeating_table(document, trigger_placeholders: list[str], row_contexts: list[dict[str, Any]]) -> bool:
    from docx.table import _Row

    for table in document.tables:
        placeholder_rows = [
            row
            for row in list(table.rows)
            if any(placeholder in row_text(row) for placeholder in trigger_placeholders)
        ]
        if not placeholder_rows:
            continue
        template_tr = placeholder_rows[0]._tr
        parent = template_tr.getparent()
        insert_at = parent.index(template_tr)
        for offset, row_context in enumerate(row_contexts):
            new_tr = deepcopy(template_tr)
            parent.insert(insert_at + offset, new_tr)
            replace_placeholders_in_row(_Row(new_tr, table), row_context)
        for placeholder_row in placeholder_rows:
            parent.remove(placeholder_row._tr)
        return True
    return False


def weekly_row_contexts(weekly_df: pd.DataFrame) -> list[dict[str, Any]]:
    rows = []
    for row in weekly_df.fillna("").to_dict("records"):
        bentuk_metode_media = "\n".join(
            item
            for item in [
                str(row.get("modalitas", "")).strip(),
                str(row.get("bentuk_pembelajaran", "")).strip(),
                str(row.get("metode", "")).strip(),
            ]
            if item
        )
        rows.append(
            {
                "minggu": row.get("minggu", ""),
                "kemampuan_akhir": clean_subcpmk_text(row.get("sub_cpmk", "")),
                "sub_cpmk": clean_subcpmk_text(row.get("sub_cpmk", "")),
                "materi": row.get("materi", ""),
                "bentuk_metode_media": bentuk_metode_media,
                "estimasi_waktu": row.get("estimasi_waktu", ""),
                "pengalaman_belajar": row.get("pengalaman_belajar", ""),
                "kriteria_bentuk": row.get("teknik_asesmen", ""),
                "indikator_penilaian": row.get("indikator_penilaian", ""),
                "bobot": f"{as_float(row.get('bobot')):g}%" if str(row.get("bobot", "")).strip() else "",
            }
        )
    return rows


def kisi_row_contexts(weekly_df: pd.DataFrame) -> list[dict[str, Any]]:
    contexts = []
    assessed_rows = [
        row
        for row in weekly_df.fillna("").to_dict("records")
        if str(row.get("teknik_asesmen", "")).strip() or as_float(row.get("bobot")) > 0
    ]
    for idx, row in enumerate(assessed_rows, start=1):
        contexts.append(
            {
                "no_kisi": idx,
                "kemampuan_akhir_kisi": clean_subcpmk_text(row.get("sub_cpmk", "")),
                "bentuk_instrumen": row.get("teknik_asesmen", ""),
                "kognitif": row.get("indikator_penilaian", ""),
                "afektif": "",
                "psikomotor": "",
                "nomor_butir_soal": row.get("minggu", ""),
            }
        )
    return contexts or [
        {
            "no_kisi": "",
            "kemampuan_akhir_kisi": "",
            "bentuk_instrumen": "",
            "kognitif": "",
            "afektif": "",
            "psikomotor": "",
            "nomor_butir_soal": "",
        }
    ]


def assessment_component_contexts(weekly_df: pd.DataFrame) -> list[dict[str, Any]]:
    totals: dict[str, float] = {}
    for row in weekly_df.fillna("").to_dict("records"):
        technique = str(row.get("teknik_asesmen", "")).strip()
        if not technique:
            continue
        totals[technique] = totals.get(technique, 0.0) + as_float(row.get("bobot"))
    return [
        {"aspek_penilaian": key, "persentase_penilaian": f"{value:g}%"}
        for key, value in totals.items()
    ] or [{"aspek_penilaian": "", "persentase_penilaian": ""}]


def evaluation_row_contexts(evaluation_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    contexts = []
    for row in evaluation_records:
        contexts.append(
            {
                "bentuk_tes": row.get("bentuk_tes", ""),
                "jenis_tes": row.get("jenis_tes", ""),
                "instrumen_penilaian": row.get("instrumen_penilaian", ""),
                "rubrik_penilaian": row.get("rubrik_penilaian", ""),
                "kisi_kisi_instrumen": row.get("kisi_kisi_instrumen", ""),
            }
        )
    return contexts or [
        {
            "bentuk_tes": "",
            "jenis_tes": "",
            "instrumen_penilaian": "",
            "rubrik_penilaian": "",
            "kisi_kisi_instrumen": "",
        }
    ]


def cpl_row_contexts(cpl_records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "kode_cpl": row.get("kode_cpl", ""),
            "rumusan_cpl": row.get("deskripsi_cpl", ""),
        }
        for row in cpl_records
    ] or [{"kode_cpl": "", "rumusan_cpl": ""}]


def cpmk_row_contexts(cpmk_df: pd.DataFrame) -> list[dict[str, Any]]:
    return [
        {
            "kode_cpmk": row.get("kode_cpmk", ""),
            "rumusan_cpmk": row.get("deskripsi_cpmk", ""),
        }
        for row in cpmk_df.fillna("").to_dict("records")
    ] or [{"kode_cpmk": "", "rumusan_cpmk": ""}]


def fill_template_tables(document, payload: dict[str, Any], cpmk_df: pd.DataFrame, weekly_df: pd.DataFrame) -> None:
    fill_repeating_table(document, ["{{kode_cpl}}", "{{rumusan_cpl}}"], cpl_row_contexts(payload["cpl"]))
    fill_repeating_table(document, ["{{kode_cpmk}}", "{{rumusan_cpmk}}"], cpmk_row_contexts(cpmk_df))
    fill_repeating_table(
        document,
        ["{{minggu}}", "{{kemampuan_akhir}}", "{{bentuk_metode_media}}"],
        weekly_row_contexts(weekly_df),
    )
    fill_repeating_table(
        document,
        ["{{aspek_penilaian}}", "{{persentase_penilaian}}"],
        assessment_component_contexts(weekly_df),
    )
    fill_repeating_table(
        document,
        ["{{bentuk_tes}}", "{{jenis_tes}}", "{{rubrik_penilaian}}"],
        evaluation_row_contexts(payload.get("evaluations", [])),
    )
    fill_repeating_table(
        document,
        ["{{no_kisi}}", "{{kemampuan_akhir_kisi}}", "{{bentuk_instrumen}}"],
        kisi_row_contexts(weekly_df),
    )


LEGACY_STRINGS = [
    "Sistem Kendali Kontinyu",
    "RTE244007",
    "Mila Fauziyah",
    "Hari Kurnia Safitri",
    "Anindya",
    "Matematika 2",
]


def selected_data_text(context: dict[str, Any], weekly_df: pd.DataFrame, cpmk_df: pd.DataFrame) -> str:
    parts = [str(value) for value in context.values()]
    parts.extend(str(value) for value in weekly_df.fillna("").astype(str).to_numpy().flatten())
    parts.extend(str(value) for value in cpmk_df.fillna("").astype(str).to_numpy().flatten())
    return "\n".join(parts)


def replace_text_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    original_text = paragraph.text
    if not original_text:
        return
    rendered_text = original_text
    for old, new in replacements.items():
        rendered_text = rendered_text.replace(old, new)
    if rendered_text == original_text:
        return
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = rendered_text
    else:
        paragraph.add_run(rendered_text)


def scrub_legacy_text(document, context: dict[str, Any], weekly_df: pd.DataFrame, cpmk_df: pd.DataFrame) -> None:
    selected_text = selected_data_text(context, weekly_df, cpmk_df)
    replacements = {
        text: "" for text in LEGACY_STRINGS if text and text not in selected_text
    }
    if not replacements:
        return
    for paragraph in document.paragraphs:
        replace_text_in_paragraph(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_paragraph(paragraph, replacements)
    for section in document.sections:
        for header_footer in [section.header, section.footer]:
            for paragraph in header_footer.paragraphs:
                replace_text_in_paragraph(paragraph, replacements)
            for table in header_footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replace_text_in_paragraph(paragraph, replacements)


def normalize_word_labels(document) -> None:
    replacements = {
        "Sub-CPMK": "Kemampuan akhir yang direncanakan",
        "Sub CPMK": "Kemampuan akhir yang direncanakan",
    }
    for paragraph in document.paragraphs:
        replace_text_in_paragraph(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_paragraph(paragraph, replacements)
    for section in document.sections:
        for header_footer in [section.header, section.footer]:
            for paragraph in header_footer.paragraphs:
                replace_text_in_paragraph(paragraph, replacements)
            for table in header_footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replace_text_in_paragraph(paragraph, replacements)


def replace_placeholders_in_paragraph(paragraph, context: dict[str, Any]) -> None:
    original_text = paragraph.text
    if "{{" not in original_text:
        return
    rendered_text = replace_text_placeholders(original_text, context)
    if rendered_text == original_text:
        return
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = rendered_text
    else:
        paragraph.add_run(rendered_text)


def replace_placeholders_in_document(document, context: dict[str, Any]) -> None:
    for paragraph in document.paragraphs:
        replace_placeholders_in_paragraph(paragraph, context)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_placeholders_in_paragraph(paragraph, context)
    for section in document.sections:
        for header_footer in [section.header, section.footer]:
            for paragraph in header_footer.paragraphs:
                replace_placeholders_in_paragraph(paragraph, context)
            for table in header_footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replace_placeholders_in_paragraph(paragraph, context)


def render_docx(
    context: dict[str, Any],
    payload: dict[str, Any],
    cpmk_df: pd.DataFrame,
    weekly_df: pd.DataFrame,
    style_messages=None,
) -> bytes:
    from docx import Document

    OUTPUT_DIR.mkdir(exist_ok=True)
    if not DEFAULT_TEMPLATE_PATH.exists():
        raise ValueError(
            "Template RPS default tidak ditemukan. Pastikan file "
            f"{DEFAULT_TEMPLATE_RELATIVE_PATH} tersedia."
        )
    template_bytes = DEFAULT_TEMPLATE_PATH.read_bytes()
    validate_template_loops(template_bytes)
    document = Document(io.BytesIO(template_bytes))
    render_context = context.copy()
    try:
        fill_template_tables(document, payload, cpmk_df, weekly_df)
        replace_placeholders_in_document(document, render_context)
        normalize_word_labels(document)
        scrub_legacy_text(document, render_context, weekly_df, cpmk_df)
    except Exception as exc:
        raise ValueError(
            "Template Word gagal diisi. Pastikan template dapat dibuka dan tidak rusak."
        ) from exc
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()


def make_excel_export(
    payload: dict[str, Any],
    lecturer_name: str,
    description: str,
    cpmk_df: pd.DataFrame,
    weekly_df: pd.DataFrame,
    reference_df: pd.DataFrame,
    validation_df: pd.DataFrame,
) -> bytes:
    buffer = io.BytesIO()
    mk = pd.DataFrame(
        [
            {
                **payload["mk"],
                "total_sks": payload["total_sks"],
                "dosen_pengampu": lecturer_name,
                "deskripsi_mk": description,
            }
        ]
    )
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        mk.to_excel(writer, sheet_name="Identitas_MK", index=False)
        pd.DataFrame(payload["cpl"]).to_excel(writer, sheet_name="CPL", index=False)
        pd.DataFrame(payload["ik"]).to_excel(writer, sheet_name="IK", index=False)
        cpmk_df.to_excel(writer, sheet_name="CPMK", index=False)
        weekly_df.to_excel(writer, sheet_name="RPS_Pertemuan", index=False)
        reference_df.to_excel(writer, sheet_name="Referensi", index=False)
        validation_df.to_excel(writer, sheet_name="Validasi", index=False)
    buffer.seek(0)
    return buffer.read()


def generate_sample_files() -> bytes:
    from docx import Document

    SAMPLE_DIR.mkdir(exist_ok=True)
    TEMPLATE_DIR.mkdir(exist_ok=True)

    def sample_cpmk_rows(kode_mk: str, nama_mk: str) -> list[dict[str, str]]:
        return [
            {
                "kode_mk": kode_mk,
                "kode_cpmk": f"{kode_mk}-CPMK-1",
                "deskripsi_cpmk": f"Mahasiswa mampu menjelaskan konsep dasar {nama_mk}.",
                "kode_ik": "IK-1.1",
            },
            {
                "kode_mk": kode_mk,
                "kode_cpmk": f"{kode_mk}-CPMK-2",
                "deskripsi_cpmk": f"Mahasiswa mampu menerapkan prosedur dan analisis pada {nama_mk}.",
                "kode_ik": "IK-2.1",
            },
            {
                "kode_mk": kode_mk,
                "kode_cpmk": f"{kode_mk}-CPMK-3",
                "deskripsi_cpmk": f"Mahasiswa mampu menyusun laporan atau produk pembelajaran {nama_mk}.",
                "kode_ik": "IK-3.1",
            },
        ]

    def sample_weekly_rows(kode_mk: str, nama_mk: str, practice: bool) -> list[dict[str, Any]]:
        mk = {
            "kode_mk": kode_mk,
            "nama_mk": nama_mk,
            "sks_praktek": 2 if practice else 0,
            "jenis_mk": "Praktik" if practice else "Teori",
        }
        cpmk = sample_cpmk_rows(kode_mk, nama_mk)
        rows = normalize_weekly_df([], mk, cpmk).to_dict("records")
        for row in rows:
            week = as_int(row["minggu"])
            if week not in [8, 16]:
                row["sub_cpmk"] = f"Menguasai topik {nama_mk} pertemuan {week}"
                row["materi"] = f"Materi {nama_mk} minggu {week}"
                row["indikator_penilaian"] = "Ketepatan konsep, prosedur, dan argumentasi"
            row["pengalaman_belajar"] = (
                "Mahasiswa melakukan praktik, mencatat hasil, dan mendiskusikan temuan."
                if practice
                else "Mahasiswa mengikuti diskusi, studi kasus, dan latihan terarah."
            )
            row["referensi"] = "Modul ajar dan sumber belajar prodi"
        return rows

    sample_data = {
        "Master_CPL": pd.DataFrame(
            [
                {
                    "kode_cpl": "CPL-1",
                    "deskripsi_cpl": "Mampu menerapkan konsep matematika, sains, dan teknologi elektro.",
                },
                {
                    "kode_cpl": "CPL-2",
                    "deskripsi_cpl": "Mampu merancang dan menganalisis sistem tenaga listrik sederhana.",
                },
                {
                    "kode_cpl": "CPL-3",
                    "deskripsi_cpl": "Mampu bekerja profesional, komunikatif, dan bertanggung jawab.",
                },
            ]
        ),
        "Master_IK": pd.DataFrame(
            [
                {
                    "kode_ik": "IK-1.1",
                    "deskripsi_ik": "Menjelaskan prinsip dasar rangkaian listrik.",
                    "kode_cpl": "CPL-1",
                },
                {
                    "kode_ik": "IK-2.1",
                    "deskripsi_ik": "Menganalisis performa komponen sistem tenaga.",
                    "kode_cpl": "CPL-2",
                },
                {
                    "kode_ik": "IK-3.1",
                    "deskripsi_ik": "Menyusun laporan teknis sesuai kaidah profesional.",
                    "kode_cpl": "CPL-3",
                },
            ]
        ),
        "Master_MK": pd.DataFrame(
            [
                {
                    "kode_mk": "PSTE2101",
                    "nama_mk": "Dasar Sistem Tenaga Listrik",
                    "nama_prodi": "D3 Program Studi Teknik Elektro",
                    "semester": 2,
                    "sks_teori": 2,
                    "sks_praktek": 1,
                    "jenis_mk": "Wajib",
                },
                {
                    "kode_mk": "PSTE3204",
                    "nama_mk": "Instalasi Listrik Industri",
                    "nama_prodi": "D3 Program Studi Teknik Elektro",
                    "semester": 3,
                    "sks_teori": 1,
                    "sks_praktek": 2,
                    "jenis_mk": "Wajib",
                },
                {
                    "kode_mk": "PSTE4301",
                    "nama_mk": "PLC",
                    "nama_prodi": "D3 Program Studi Teknik Elektro",
                    "semester": 4,
                    "sks_teori": 1,
                    "sks_praktek": 2,
                    "jenis_mk": "Praktik",
                },
                {
                    "kode_mk": "PSTE4302",
                    "nama_mk": "Instrumentasi Industri",
                    "nama_prodi": "D3 Program Studi Teknik Elektro",
                    "semester": 4,
                    "sks_teori": 2,
                    "sks_praktek": 1,
                    "jenis_mk": "Wajib",
                },
                {
                    "kode_mk": "PSTE6401",
                    "nama_mk": "Proyek Akhir",
                    "nama_prodi": "D3 Program Studi Teknik Elektro",
                    "semester": 6,
                    "sks_teori": 0,
                    "sks_praktek": 4,
                    "jenis_mk": "Proyek",
                },
            ]
        ),
        "Mapping_MK_CPL": pd.DataFrame(
            [
                {"kode_mk": "PSTE2101", "kode_cpl": "CPL-1"},
                {"kode_mk": "PSTE2101", "kode_cpl": "CPL-2"},
                {"kode_mk": "PSTE2101", "kode_cpl": "CPL-3"},
                {"kode_mk": "PSTE3204", "kode_cpl": "CPL-2"},
                {"kode_mk": "PSTE3204", "kode_cpl": "CPL-3"},
                {"kode_mk": "PSTE4301", "kode_cpl": "CPL-1"},
                {"kode_mk": "PSTE4301", "kode_cpl": "CPL-2"},
                {"kode_mk": "PSTE4301", "kode_cpl": "CPL-3"},
                {"kode_mk": "PSTE4302", "kode_cpl": "CPL-1"},
                {"kode_mk": "PSTE4302", "kode_cpl": "CPL-2"},
                {"kode_mk": "PSTE4302", "kode_cpl": "CPL-3"},
                {"kode_mk": "PSTE6401", "kode_cpl": "CPL-1"},
                {"kode_mk": "PSTE6401", "kode_cpl": "CPL-2"},
                {"kode_mk": "PSTE6401", "kode_cpl": "CPL-3"},
            ]
        ),
        "Master_CPMK": pd.DataFrame(
            [
                *sample_cpmk_rows("PSTE2101", "Dasar Sistem Tenaga Listrik"),
                *sample_cpmk_rows("PSTE4301", "PLC"),
                *sample_cpmk_rows("PSTE4302", "Instrumentasi Industri"),
                *sample_cpmk_rows("PSTE6401", "Proyek Akhir"),
            ]
        ),
        "Short_Silabus": pd.DataFrame(
            [
                {
                    "kode_mk": "PSTE2101",
                    "deskripsi_mk": "Mata kuliah ini membahas konsep dasar pembangkitan, transmisi, distribusi, dan beban listrik.",
                    "bahan_kajian": "Konsep sistem tenaga; pembangkit; transmisi; distribusi; proteksi dasar; analisis beban.",
                },
                {
                    "kode_mk": "PSTE4301",
                    "deskripsi_mk": "Mata kuliah PLC membahas prinsip kontrol logika terprogram, pemrograman ladder, wiring I/O, dan penerapan otomasi industri.",
                    "bahan_kajian": "Arsitektur PLC; input-output digital; ladder diagram; timer-counter; interlock; troubleshooting; mini project otomasi.",
                },
                {
                    "kode_mk": "PSTE4302",
                    "deskripsi_mk": "Mata kuliah Instrumentasi Industri membahas sensor, transduser, pengukuran proses, kalibrasi, dan integrasi instrumentasi pada sistem industri.",
                    "bahan_kajian": "Sensor industri; transduser; pengkondisi sinyal; kalibrasi; aktuator; loop kontrol; dokumentasi instrumentasi.",
                },
                {
                    "kode_mk": "PSTE6401",
                    "deskripsi_mk": "Mata kuliah Proyek Akhir membimbing mahasiswa merancang, membangun, menguji, dan melaporkan solusi teknologi terapan sesuai bidang elektro.",
                    "bahan_kajian": "Perencanaan proyek; desain teknis; implementasi; pengujian; analisis hasil; laporan akhir; presentasi dan demonstrasi.",
                },
            ]
        ),
        "RPS_Pertemuan": pd.DataFrame(
            sample_weekly_rows("PSTE2101", "Dasar Sistem Tenaga Listrik", False)
            + sample_weekly_rows("PSTE4301", "PLC", True)
            + sample_weekly_rows("PSTE4302", "Instrumentasi Industri", True)
            + sample_weekly_rows("PSTE6401", "Proyek Akhir", True)
        ),
        "Referensi": pd.DataFrame(
            [
                {
                    "kode_mk": "PSTE2101",
                    "referensi": "Grainger, J. J., & Stevenson, W. D. Power System Analysis.",
                },
                {
                    "kode_mk": "PSTE2101",
                    "referensi": "Glover, J. D., Sarma, M. S., & Overbye, T. Power System Analysis and Design.",
                },
                {
                    "kode_mk": "PSTE2101",
                    "referensi": "Materi ajar Program Studi Teknik Elektro.",
                },
                {
                    "kode_mk": "PSTE4301",
                    "referensi": "Petruzella, F. D. Programmable Logic Controllers.",
                },
                {
                    "kode_mk": "PSTE4301",
                    "referensi": "Manual PLC dan modul praktikum otomasi industri prodi.",
                },
                {
                    "kode_mk": "PSTE4302",
                    "referensi": "Doebelin, E. O. Measurement Systems: Application and Design.",
                },
                {
                    "kode_mk": "PSTE4302",
                    "referensi": "Modul sensor, transduser, dan instrumentasi industri prodi.",
                },
                {
                    "kode_mk": "PSTE6401",
                    "referensi": "Panduan Proyek Akhir Program Studi Teknik Elektro.",
                },
                {
                    "kode_mk": "PSTE6401",
                    "referensi": "Standar penulisan laporan dan presentasi karya teknologi terapan.",
                },
            ]
        ),
    }

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        for sheet_name, df in sample_data.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    excel_bytes = excel_buffer.getvalue()
    DEFAULT_SAMPLE_PATH.write_bytes(excel_bytes)
    return excel_bytes

    doc = Document()
    doc.add_heading("RENCANA PEMBELAJARAN SEMESTER", level=1)
    doc.add_paragraph("Program Studi: {{ nama_prodi }}")
    doc.add_paragraph("Kode MK: {{ kode_mk }}")
    doc.add_paragraph("Nama MK: {{ nama_mk }}")
    doc.add_paragraph("Semester: {{ semester }}")
    doc.add_paragraph("SKS Teori: {{ sks_teori }}")
    doc.add_paragraph("SKS Praktek: {{ sks_praktek }}")
    doc.add_paragraph("Total SKS: {{ total_sks }}")
    doc.add_paragraph("Dosen Pengampu: {{ dosen_pengampu }}")
    doc.add_heading("Deskripsi Mata Kuliah", level=2)
    doc.add_paragraph("{{ deskripsi_mk }}")
    doc.add_heading("CPL yang Dibebankan", level=2)
    doc.add_paragraph("{{ cpl_text }}")
    doc.add_heading("IK Terkait", level=2)
    doc.add_paragraph("{{ ik_text }}")
    doc.add_heading("CPMK", level=2)
    doc.add_paragraph("{{ cpmk_text }}")
    doc.add_heading("Bahan Kajian", level=2)
    doc.add_paragraph("{{ bahan_kajian }}")
    doc.add_heading("Rencana Mingguan", level=2)
    doc.add_paragraph("{{ rencana_mingguan_text }}")
    doc.add_heading("Asesmen", level=2)
    doc.add_paragraph("{{ asesmen_text }}")
    doc.add_heading("Referensi", level=2)
    doc.add_paragraph("{{ referensi_text }}")
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    template_bytes = doc_buffer.getvalue()
    DEFAULT_TEMPLATE_PATH.write_bytes(template_bytes)

    return excel_bytes


def load_or_create_sample_files() -> bytes:
    if DEFAULT_SAMPLE_PATH.exists():
        return DEFAULT_SAMPLE_PATH.read_bytes()
    return generate_sample_files()


def load_program_sample(program: str) -> tuple[bytes, str]:
    if program == "D4 Teknik Elektronika":
        if not DEFAULT_D4_SAMPLE_PATH.exists():
            raise FileNotFoundError("Master D4 belum tersedia di sample_data.")
        return DEFAULT_D4_SAMPLE_PATH.read_bytes(), DEFAULT_D4_SAMPLE_PATH.name
    return load_or_create_sample_files(), DEFAULT_D3_SAMPLE_PATH.name


def show_preview(payload: dict[str, Any]) -> None:
    mk = payload["mk"]
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Kode MK", mk.get("kode_mk", "-"))
    col2.metric("Semester", mk.get("semester", "-"))
    col3.metric("Total SKS", payload["total_sks"])
    col4.metric("Jenis MK", mk.get("jenis_mk", "-"))
    total_hours = as_float(mk.get("total_jam_minggu"))
    if total_hours:
        st.caption(
            f"Beban per minggu: {as_float(mk.get('jam_teori_minggu')):g} jam teori + "
            f"{as_float(mk.get('jam_praktik_minggu')):g} jam praktik = {total_hours:g} jam."
        )
    with st.expander("Short silabus", expanded=True):
        st.write(payload["silabus"].get("deskripsi_mk", "-"))
        st.caption(payload["silabus"].get("bahan_kajian", ""))
    for warning in payload.get("warnings", []):
        st.warning(warning)


def main() -> None:
    st.set_page_config(page_title="RPS Builder OBE", layout="wide")
    st.title("RPS Builder OBE")
    st.caption("Mengisi template Word RPS prodi dari master kurikulum Excel.")

    OUTPUT_DIR.mkdir(exist_ok=True)
    with st.sidebar:
        st.header("Input")
        program = st.selectbox(
            "Jenjang / master kurikulum",
            ["D3 Teknik Elektro", "D4 Teknik Elektronika"],
            help="Master D3 dan D4 disimpan terpisah agar data tidak saling menggantikan.",
        )
        try:
            sample_excel, sample_filename = load_program_sample(program)
        except FileNotFoundError as exc:
            st.error(str(exc))
            st.stop()
        excel_file = st.file_uploader(
            "Upload Excel master kurikulum",
            type=["xlsx"],
            help="Gunakan sheet sesuai format master kurikulum prodi.",
            key=f"master_upload_{program}",
        )
        st.download_button(
            f"Unduh master {program.split()[0]}",
            data=sample_excel,
            file_name=sample_filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    excel_bytes = excel_file.getvalue() if excel_file else sample_excel

    try:
        workbook = load_master_excel(excel_bytes)
    except Exception as exc:
        st.error(f"Excel tidak dapat dibaca: {exc}")
        st.stop()

    schema_errors = validate_workbook_schema(workbook)
    if schema_errors:
        st.error("Format Excel belum sesuai.")
        for error in schema_errors:
            st.write(f"- {error}")
        st.stop()

    master_mk = workbook["Master_MK"].copy()
    master_mk["kode_mk"] = master_mk["kode_mk"].map(normalize_kode_mk)
    master_mk = master_mk[master_mk["kode_mk"] != ""].copy()
    master_mk["course_key"] = dataframe_course_keys(master_mk)
    duplicate_codes = set(
        master_mk.loc[master_mk["kode_mk"].duplicated(keep=False), "kode_mk"]
    )
    labels: dict[str, str] = {}
    for row in master_mk.fillna("").to_dict("records"):
        key = course_key_from_row(row)
        label = f"{row['kode_mk']} - {row['nama_mk']} (Semester {row['semester']})"
        if row["kode_mk"] in duplicate_codes:
            label += f" [{key}]"
        labels[key] = label
    selected_key = st.selectbox(
        "Pilih mata kuliah", master_mk["course_key"].tolist(), format_func=labels.get
    )
    payload = build_course_payload(workbook, selected_key)
    selected_code = normalize_kode_mk(payload["mk"].get("kode_mk", ""))
    selected_offer = normalize_id_penawaran(payload["mk"].get("id_penawaran", ""))
    editor_key = re.sub(r"[^A-Za-z0-9_-]+", "_", selected_key)

    main_tabs = st.tabs(
        [
            "Identitas MK",
            "CPL, IK, dan CPMK",
            "RPS Pertemuan",
            "Preview Export Word",
            "Validasi Data RPS",
        ]
    )

    with main_tabs[0]:
        st.subheader(payload["mk"].get("nama_mk", "Mata Kuliah"))
        show_preview(payload)
        lecturer_name = st.text_input("Nama dosen pengampu", value="")
        description = st.text_area(
            "Deskripsi mata kuliah",
            value=str(payload["silabus"].get("deskripsi_mk", "")),
            height=120,
        )
        reference_df = st.data_editor(
            records_to_editor(payload["references"], ["kode_mk", "referensi"]),
            use_container_width=True,
            num_rows="dynamic",
            key=f"reference_{editor_key}",
        )
        reference_df["kode_mk"] = selected_code
        reference_df["id_penawaran"] = selected_offer

    with main_tabs[1]:
        st.subheader("CPL yang dibebankan pada MK")
        st.dataframe(pd.DataFrame(payload["cpl"]), use_container_width=True, hide_index=True)
        st.subheader("IK terkait")
        st.dataframe(pd.DataFrame(payload["ik"]), use_container_width=True, hide_index=True)
        st.subheader("CPMK")
        cpmk_df = st.data_editor(
            records_to_editor(
                payload["cpmk"],
                ["kode_mk", "kode_cpmk", "deskripsi_cpmk", "kode_ik"],
            ),
            use_container_width=True,
            num_rows="dynamic",
            key=f"cpmk_{editor_key}",
        )
        cpmk_df["kode_mk"] = selected_code
        cpmk_df["id_penawaran"] = selected_offer

    with main_tabs[2]:
        assessment_options = unique_values(
            [""] + TEKNIK_ASESMEN_OPTIONS + [
                str(row.get("teknik_asesmen", "")).strip()
                for row in payload["weekly"]
                if str(row.get("teknik_asesmen", "")).strip()
            ]
        )
        weekly_df = st.data_editor(
            records_to_editor(payload["weekly"], RPS_WEEKLY_COLUMNS),
            use_container_width=True,
            num_rows="fixed",
            height=620,
            disabled=["kode_mk", "id_penawaran", "minggu"],
            column_order=[
                "minggu",
                "sub_cpmk",
                "materi",
                "modalitas",
                "bentuk_pembelajaran",
                "metode",
                "pengalaman_belajar",
                "teknik_asesmen",
                "indikator_penilaian",
                "bobot",
                "referensi",
                "kode_cpmk",
            ],
            column_config={
                "minggu": st.column_config.NumberColumn(
                    RPS_WEEKLY_LABELS["minggu"], min_value=1, max_value=17, step=1
                ),
                "sub_cpmk": st.column_config.TextColumn(
                    RPS_WEEKLY_LABELS["sub_cpmk"], width="large"
                ),
                "materi": st.column_config.TextColumn(
                    RPS_WEEKLY_LABELS["materi"], width="large"
                ),
                "modalitas": st.column_config.SelectboxColumn(
                    RPS_WEEKLY_LABELS["modalitas"], options=MODALITAS_OPTIONS
                ),
                "bentuk_pembelajaran": st.column_config.SelectboxColumn(
                    RPS_WEEKLY_LABELS["bentuk_pembelajaran"], options=BENTUK_OPTIONS
                ),
                "metode": st.column_config.SelectboxColumn(
                    RPS_WEEKLY_LABELS["metode"], options=METODE_OPTIONS
                ),
                "pengalaman_belajar": st.column_config.TextColumn(
                    RPS_WEEKLY_LABELS["pengalaman_belajar"], width="large"
                ),
                "teknik_asesmen": st.column_config.SelectboxColumn(
                    RPS_WEEKLY_LABELS["teknik_asesmen"], options=assessment_options
                ),
                "indikator_penilaian": st.column_config.TextColumn(
                    RPS_WEEKLY_LABELS["indikator_penilaian"], width="large"
                ),
                "bobot": st.column_config.NumberColumn(
                    RPS_WEEKLY_LABELS["bobot"], min_value=0.0, max_value=100.0, step=1.0
                ),
                "referensi": st.column_config.TextColumn(
                    RPS_WEEKLY_LABELS["referensi"], width="medium"
                ),
            },
            key=f"weekly_{editor_key}",
        )
        weekly_df["kode_mk"] = selected_code
        weekly_df["id_penawaran"] = selected_offer
        weekly_df = weekly_df[RPS_WEEKLY_COLUMNS]
        total_weight = weekly_df.get("bobot", pd.Series(dtype=float)).map(as_float).sum()
        st.metric("Total Bobot Penilaian", f"{total_weight:g}%")
        if round(total_weight, 2) != 100:
            st.warning("Total bobot penilaian belum 100%.")

    validation_df = validate_rps(payload, cpmk_df, weekly_df)
    master_validation_df = validate_master_data(workbook)
    has_error = (validation_df["status"] == "Error").any()

    context = make_context(
        payload,
        lecturer_name,
        description,
        cpmk_df,
        weekly_df,
        reference_df,
    )
    docx_bytes = None
    docx_error = ""
    table_style_messages: list[str] = []
    try:
        docx_bytes = render_docx(context, payload, cpmk_df, weekly_df, table_style_messages)
    except ValueError as exc:
        docx_error = str(exc)
    excel_export = make_excel_export(
        payload,
        lecturer_name,
        description,
        cpmk_df,
        weekly_df,
        reference_df,
        validation_df,
    )
    validation_export = validation_df.to_csv(index=False).encode("utf-8")
    filename_base = f"RPS_{selected_code}_{str(payload['mk'].get('nama_mk', 'MK')).replace(' ', '_')}"

    with main_tabs[3]:
        st.subheader("Preview Export Word")
        st.write(f"Template: `{DEFAULT_TEMPLATE_RELATIVE_PATH}`")
        st.write(f"Mata kuliah: `{payload['mk'].get('kode_mk', '')} - {payload['mk'].get('nama_mk', '')}`")
        st.write("Tabel RPS pertemuan akan diisi pada posisi tabel yang sudah ada di template.")
        if docx_bytes:
            for message in dict.fromkeys(table_style_messages):
                st.warning(message)
            col1, col2, col3 = st.columns(3)
            col1.download_button(
                "Download RPS Word",
                data=docx_bytes,
                file_name=f"{filename_base}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            col2.download_button(
                "Download RPS Excel",
                data=excel_export,
                file_name=f"{filename_base}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            col3.download_button(
                "Download laporan validasi",
                data=validation_export,
                file_name=f"validasi_{selected_code}.csv",
                mime="text/csv",
            )
        else:
            st.error(docx_error)

    with main_tabs[4]:
        st.subheader("Validasi RPS terpilih")
        st.dataframe(validation_df, use_container_width=True, hide_index=True)
        if has_error:
            st.warning("Masih ada validasi berstatus Error. Export tetap bisa dibuat untuk draft.")
        else:
            st.success("Validasi utama terpenuhi.")
        st.subheader("Validasi data master RPS")
        st.dataframe(master_validation_df, use_container_width=True, hide_index=True)


if __name__ == "__main__":
    main()
